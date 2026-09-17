from __future__ import annotations

import json
import math
import statistics
import sys
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


MONTHS = {
    "JANEIRO", "FEVEREIRO", "MARÇO", "ABRIL", "MAIO", "JUNHO",
    "JULHO", "AGOSTO", "SETEMBRO", "OUTUBRO", "NOVEMBRO", "DEZEMBRO",
}


def clean(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and math.isnan(value):
        return None
    return value


def num(value):
    return float(value) if isinstance(value, (int, float)) and not isinstance(value, bool) else None


def mean(values):
    values = [value for value in values if value is not None]
    return statistics.fmean(values) if values else None


def median(values):
    values = [value for value in values if value is not None]
    return statistics.median(values) if values else None


def safe_rate(numerator, denominator):
    if numerator is None or denominator in (None, 0):
        return None
    return numerator / denominator


workbook_path = Path(sys.argv[1])
wb = load_workbook(workbook_path, data_only=True, read_only=False)
ws = wb["Postagens 2026"]

month_markers = []
for row in range(1, ws.max_row + 1):
    value = ws.cell(row, 1).value
    if isinstance(value, str) and value.strip().upper() in MONTHS:
        month_markers.append((row, value.strip().upper()))

september_row = next(row for row, label in month_markers if label == "SETEMBRO")
next_rows = [row for row, _ in month_markers if row > september_row]
september_end = (min(next_rows) - 1) if next_rows else ws.max_row

header_hits = []
keywords = ("visual", "alcance", "reten", "gancho", "intera", "compart", "coment", "salv", "seguidor")
for row in range(1, min(ws.max_row, 200) + 1):
    for col in range(1, min(ws.max_column, 20) + 1):
        value = ws.cell(row, col).value
        if isinstance(value, str) and any(k in value.lower() for k in keywords):
            header_hits.append({"cell": ws.cell(row, col).coordinate, "value": value})

posts = []
for row in range(september_row + 1, september_end + 1):
    title = ws.cell(row, 1).value
    fmt = ws.cell(row, 2).value
    posted_at = ws.cell(row, 3).value
    period = ws.cell(row, 4).value
    if not (
        isinstance(title, str)
        and title.strip()
        and isinstance(fmt, str)
        and fmt.strip()
        and isinstance(posted_at, (datetime, date))
        and isinstance(period, str)
        and period.strip().lower().startswith("24h")
    ):
        continue

    snapshots = []
    for snapshot_row in range(row, min(row + 4, september_end + 1)):
        snapshot_title = ws.cell(snapshot_row, 1).value
        if snapshot_title != title:
            break
        phase_value = ws.cell(snapshot_row, 3).value
        label_value = ws.cell(snapshot_row, 4).value
        phase = "24h" if snapshot_row == row else str(phase_value or label_value or "").strip()
        values = [clean(ws.cell(snapshot_row, col).value) for col in range(1, 17)]
        snapshots.append({
            "row": snapshot_row,
            "phase": phase,
            "views": num(values[4]),
            "reach": num(values[5]),
            "retention_seconds": num(values[6]),
            "hook_rate": num(values[7]),
            "interactions": num(values[8]),
            "interaction_rate": num(values[9]),
            "shares": num(values[10]),
            "share_rate_reported": num(values[11]),
            "comments": num(values[12]),
            "saves": num(values[13]),
            "followers": num(values[14]),
            "note": values[15],
        })

    baseline = snapshots[0]
    posted_date = posted_at.date() if isinstance(posted_at, datetime) else posted_at
    posts.append({
        "row": row,
        "title": title.strip(),
        "format": fmt.strip(),
        "date": posted_date.isoformat(),
        "weekday": posted_date.strftime("%A"),
        "slot": "Dia" if "dia" in period.lower() else "Noite" if "noite" in period.lower() else None,
        "views_24h": baseline["views"],
        "reach_24h": baseline["reach"],
        "interactions_24h": baseline["interactions"],
        "shares_24h": baseline["shares"],
        "share_rate_reach_24h": safe_rate(baseline["shares"], baseline["reach"]),
        "share_rate_views_24h": safe_rate(baseline["shares"], baseline["views"]),
        "reported_share_rate_24h": baseline["share_rate_reported"],
        "comments_24h": baseline["comments"],
        "saves_24h": baseline["saves"],
        "followers_24h": baseline["followers"],
        "retention_seconds_24h": baseline["retention_seconds"],
        "hook_rate_24h": baseline["hook_rate"],
        "note_24h": baseline["note"],
        "snapshots": snapshots,
    })


def aggregate(items):
    return {
        "posts": len(items),
        "shares_total": sum((p["shares_24h"] or 0) for p in items),
        "shares_mean": mean([p["shares_24h"] for p in items]),
        "shares_median": median([p["shares_24h"] for p in items]),
        "share_rate_reach_mean": mean([p["share_rate_reach_24h"] for p in items]),
        "views_mean": mean([p["views_24h"] for p in items]),
        "reach_mean": mean([p["reach_24h"] for p in items]),
        "interactions_mean": mean([p["interactions_24h"] for p in items]),
        "saves_mean": mean([p["saves_24h"] for p in items]),
        "hook_rate_mean": mean([p["hook_rate_24h"] for p in items]),
        "retention_mean": mean([p["retention_seconds_24h"] for p in items]),
    }


by_format = defaultdict(list)
by_slot = defaultdict(list)
by_weekday = defaultdict(list)
for post in posts:
    by_format[post["format"]].append(post)
    by_slot[post["slot"]].append(post)
    by_weekday[post["weekday"]].append(post)

valid_posts = [post for post in posts if post["views_24h"] is not None]
valid_by_format = defaultdict(list)
for post in valid_posts:
    valid_by_format[post["format"]].append(post)

total_reach = sum((post["reach_24h"] or 0) for post in valid_posts)
total_shares = sum((post["shares_24h"] or 0) for post in valid_posts)
ranked_shares = sorted(valid_posts, key=lambda p: (p["shares_24h"] or 0), reverse=True)

payload = {
    "workbook": str(workbook_path),
    "sheet": ws.title,
    "dimensions": {"rows": ws.max_row, "columns": ws.max_column},
    "month_markers": month_markers,
    "september_range": f"A{september_row}:P{september_end}",
    "header_hits": header_hits,
    "posts": posts,
    "overall": aggregate(posts),
    "by_format": {key: aggregate(value) for key, value in by_format.items()},
    "by_slot": {str(key): aggregate(value) for key, value in by_slot.items()},
    "by_weekday": {key: aggregate(value) for key, value in by_weekday.items()},
    "top_by_shares": sorted(posts, key=lambda p: (p["shares_24h"] or -1), reverse=True),
    "top_by_share_rate": sorted(posts, key=lambda p: (p["share_rate_reach_24h"] or -1), reverse=True),
}

if len(sys.argv) > 2 and sys.argv[2] == "summary":
    compact_posts = []
    for post in valid_posts:
        compact_posts.append({
            key: post[key]
            for key in (
                "row", "title", "format", "date", "views_24h", "reach_24h",
                "interactions_24h", "shares_24h", "share_rate_reach_24h",
                "comments_24h", "saves_24h", "followers_24h",
                "retention_seconds_24h", "hook_rate_24h", "note_24h",
            )
        })
    summary = {
        "sheet": ws.title,
        "september_range": f"A{september_row}:P{september_end}",
        "header_hits": header_hits,
        "posts_with_24h_data": len(valid_posts),
        "posts_without_24h_data": [
            {"row": p["row"], "title": p["title"], "date": p["date"], "note": p["note_24h"]}
            for p in posts if p["views_24h"] is None
        ],
        "totals": {
            "views": sum((p["views_24h"] or 0) for p in valid_posts),
            "reach": total_reach,
            "interactions": sum((p["interactions_24h"] or 0) for p in valid_posts),
            "shares": total_shares,
            "comments": sum((p["comments_24h"] or 0) for p in valid_posts),
            "saves": sum((p["saves_24h"] or 0) for p in valid_posts),
            "followers": sum((p["followers_24h"] or 0) for p in valid_posts),
            "weighted_share_rate_by_reach": safe_rate(total_shares, total_reach),
        },
        "top_share_concentration": safe_rate(
            ranked_shares[0]["shares_24h"] if ranked_shares else None,
            total_shares,
        ),
        "by_format": {key: aggregate(value) for key, value in valid_by_format.items()},
        "posts": compact_posts,
    }
    print(json.dumps(summary, ensure_ascii=True, indent=2))
else:
    print(json.dumps(payload, ensure_ascii=True, indent=2))
